from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\USER\JuegoBoceto\Topo-v.2")
TEMPLATE = Path(r"C:\Users\USER\Downloads\PVID-420_FORMATOALUMNOTRABAJOFINAL.docx")
MEDIA = ROOT / ".docx_review" / "images_package" / "word" / "media"
OUT = ROOT / "Trabajo_Final_AR_Golpea_al_Topo.docx"
copy2(TEMPLATE, OUT)
doc = Document(OUT)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell(cell, text, size=9, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def append_cell(cell, text, size=9, bold=False):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    return p

def add_picture(cell, filename, caption):
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(MEDIA / filename), width=Inches(3.25))
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(caption); r.italic=True; r.font.size=Pt(8)

# Datos del estudiante: no se inventan los datos institucionales.
t = doc.tables[0]
set_cell(t.cell(0,1), '[POR COMPLETAR: APELLIDOS Y NOMBRES]')
set_cell(t.cell(0,4), '[POR COMPLETAR: ID]')
set_cell(t.cell(1,1), '[POR COMPLETAR: DIRECCION ZONAL O CFP]')
set_cell(t.cell(2,1), '[POR COMPLETAR: CARRERA]')
set_cell(t.cell(2,4), '[POR COMPLETAR: SEMESTRE]')
set_cell(t.cell(3,1), 'Desarrollo de aplicaciones de realidad aumentada')
set_cell(t.cell(4,1), 'Videojuego de realidad aumentada Golpea al Topo')

# Problemática y propuesta en los apartados narrativos de la plantilla.
texts = {
    'Identifica la problemática del caso práctico propuesto.':
        'La problemática consiste en desarrollar una experiencia lúdica de realidad aumentada que funcione de manera estable sobre una superficie real y permita al usuario comprender con claridad dónde se jugará, cuándo puede empezar y cómo interactuar. El reto no es solo colocar modelos 3D: la zona debe mantenerse dentro de la superficie detectada, los objetivos deben aparecer sin superponerse y los toques destinados a los controles de interfaz no deben confundirse con golpes dentro de la escena.',
    'Identifica propuesta de solución y evidencias.':
        'Se propone el videojuego AR Golpea al Topo. El usuario explora el suelo o una mesa con la cámara, toca una superficie horizontal para fijar una zona circular y pulsa Jugar. Durante una partida de 60 segundos aparecen hasta seis hoyos en posiciones separadas dentro de la zona; los topos se golpean con un toque directo. La aplicación muestra puntos, combo y tiempo, y termina con una pantalla de resultado. Las capturas del proyecto evidencian la jerarquía AR, los componentes de colocación, el gestor de partida, el controlador del martillo y las pantallas de inicio, juego y cierre.',
    'Describe la propuesta determinada para la solución del caso práctico':
        'La solución integra detección de plano horizontal, selección de una zona de juego, generación controlada de objetivos y respuesta táctil. La aplicación busca un plano dentro de su polígono y, una vez que el usuario define la zona, deja de aceptar nuevos puntos para evitar movimientos accidentales durante la partida. Cada topo aparece en un hoyo temporal dentro del área válida; si recibe un toque, se registra el puntaje y se muestra la animación de golpe. Si no se golpea a tiempo, desaparece y la racha se reinicia. El ciclo mantiene una mecánica clara y visible para el usuario.',
    'Resolver el caso práctico, utilizando como referencia el problema propuesto y las preguntas guía proporcionadas para orientar el desarrollo.':
        'La ejecución se organiza desde la preparación del entorno AR hasta las pruebas de interacción. Se priorizó que la zona de juego sea visible antes de iniciar, que los objetivos no aparezcan fuera de la superficie y que la interfaz comunique el estado de la partida. La propuesta responde a las preguntas guía mediante decisiones aplicadas al proyecto, no como una explicación aislada de herramientas.',
    'Fundamentar sus propuestas en los conocimientos adquiridos a lo largo del curso, aplicando lo aprendido en las tareas y operaciones descritas en los contenidos curriculares.':
        'La propuesta aplica programación orientada a eventos, detección espacial con AR Foundation, rayos desde la pantalla para seleccionar objetos, control de estados de juego, interfaz de usuario y pruebas funcionales. El diseño diferencia la fase de escaneo, la fijación de zona, la partida y el cierre, lo que reduce acciones inválidas y facilita el mantenimiento del flujo.',
    'Verificar el cumplimiento de los procesos desarrollados en la propuesta de solución del caso práctico.':
        'El control se realiza comprobando que la aplicación detecte una superficie horizontal, habilite el botón de inicio solo después de fijar la zona, genere objetivos dentro de los límites, registre únicamente los golpes sobre topos y actualice puntaje, combo y temporizador. También se verifica que la pantalla final muestre el resultado y permita reiniciar o cambiar de zona.',
    'Califica el impacto que representa la propuesta de solución ante la situación planteada en el caso práctico.':
        'La propuesta es viable como experiencia móvil de corta duración porque utiliza interacción por toque, señales visuales claras y una regla de juego fácil de aprender. Su principal aporte es convertir una superficie cotidiana en un espacio de juego delimitado, manteniendo la relación entre el entorno físico y los objetos virtuales. La viabilidad final en distintos dispositivos debe validarse con pruebas en equipos compatibles con ARCore o ARKit.'
}
for p in doc.paragraphs:
    key = p.text.strip()
    if key in texts:
        p.add_run('\n' + texts[key])

# Preguntas guía, sustituyendo las respuestas preliminares por respuestas aplicadas al proyecto.
q = doc.tables[1]
answers = [
    'La detección se restringe a planos horizontales y a puntos que caen dentro del polígono de la superficie reconocida. Al fijar la zona se muestra un marcador circular verde; así el usuario confirma el lugar antes de iniciar. Para las pruebas se recomienda iluminación uniforme, movimiento lento del móvil y superficies con textura, evitando reflejos intensos.',
    'En esta propuesta la navegación se resuelve como selección de zona, no como desplazamiento libre del avatar: el usuario toca una superficie detectada y puede tocar otro punto para recolocar la zona antes de empezar. Durante la partida, los objetivos se redistribuyen dentro de ese límite, por lo que la actividad permanece accesible desde la posición física del usuario.',
    'La interacción se basa en un toque que proyecta un rayo desde la cámara hacia la capa de los topos. Cuando el rayo impacta un objetivo visible, se registra un único golpe, se incrementa la puntuación, se actualiza el combo y aparece un martillo animado en el punto de contacto. Los toques sobre botones se filtran para que no activen golpes por error.',
    'La iluminación debe permitir a la cámara reconocer la superficie y, a la vez, evitar que los elementos virtuales pierdan legibilidad. Se consideran intensidad, dirección, color y sombras. En una versión de despliegue se puede activar estimación de luz del entorno; en las pruebas se debe evitar contraluz, reflejos y cambios bruscos que afecten el seguimiento.',
    'Los controles se simplifican a gestos cotidianos: tocar para fijar la zona y tocar para golpear. Los estados se comunican mediante mensajes, marcador circular, botón Jugar, marcador de puntos, combo y tiempo. La interfaz evita que un toque sobre un botón se interprete como acción sobre un objeto del juego.'
]
for row, answer in zip([1,3,5,7,9], answers): set_cell(q.cell(row,1), answer, 9)

# Cronograma de seis actividades.
cron = doc.tables[2]
activities = [
    'Analisis del caso y definicion de mecanicas', 'Configuracion de Unity y entorno AR',
    'Zona de juego y deteccion de superficies', 'Objetivos, golpes y puntaje',
    'Interfaz, pruebas y correcciones', 'Documentacion y presentacion final'
]
for i, act in enumerate(activities, 1):
    set_cell(cron.cell(i+1,0), str(i), 8)
    set_cell(cron.cell(i+1,1), act, 8)
    for c in range(2,8): set_cell(cron.cell(i+1,c), 'X' if c in ({2,3} if i==1 else {3,4} if i in (2,3) else {4,5} if i==4 else {5,6} if i==5 else {6,7}) else '', 8)

# Recursos.
for table, entries in [
    (doc.tables[3], [('Computadora para desarrollo', '1'), ('Telefono movil compatible con AR', '1'), ('Camara del dispositivo movil', '1'), ('Conexion de datos o Wi-Fi para pruebas', '1')]),
    (doc.tables[4], [('Unity Hub y Unity Editor', '1'), ('AR Foundation y proveedores AR', '1'), ('Editor de codigo', '1'), ('Herramientas de prueba y depuracion', '1')]),
    (doc.tables[5], [('Modelos 3D de topo, hoyo y martillo', '1 conjunto'), ('Materiales y texturas del juego', '1 conjunto'), ('Capturas de evidencia', '1 documento'), ('Bateria o cargador para movil', '1')])]:
    for r, (name, qty) in enumerate(entries, 2): set_cell(table.cell(r,0), name, 8); set_cell(table.cell(r,1), qty, 8)

# Propuesta y evidencias concretas.
proposal = doc.tables[6].cell(1,0)
set_cell(proposal, 'El videojuego AR Golpea al Topo transforma una superficie horizontal detectada en una zona de juego. Al iniciar, el sistema espera que el usuario explore el entorno. Luego permite fijar un circulo de juego mediante un toque y habilita el inicio. Durante los 60 segundos de partida se generan hasta seis hoyos separados entre si y dentro de la zona valida. Los topos comunes, amarillos y rojos otorgan puntajes diferentes; los aciertos consecutivos generan un multiplicador. La experiencia finaliza mostrando puntuacion y mejor combo.', 9)
for fn, cap in [('image3.png','Figura 1. Jerarquia AR y elementos principales de la escena.'), ('image14.png','Figura 2. Zona de juego fijada sobre una superficie detectada.'), ('image13.png','Figura 3. Objetivos virtuales distribuidos en la zona.'), ('image15.png','Figura 4. Indicadores de puntos, combo y tiempo.'), ('image16.png','Figura 5. Pantalla de resultado y opciones posteriores.')]: add_picture(proposal, fn, cap)

# Operaciones, normas y controles.
ops = doc.tables[7]
steps = [
('1. Analizar el caso y definir el flujo: escaneo, zona, juego y cierre.', 'Aplicar una secuencia de estados clara; registrar requisitos antes de programar.'),
('2. Configurar el proyecto para realidad aumentada en dispositivo movil.', 'Usar versiones compatibles de Unity, AR Foundation y proveedor AR; mantener copias de seguridad.'),
('3. Preparar una escena con XR Origin, camara AR y gestor de planos.', 'Comprobar permisos de camara y probar en un dispositivo compatible.'),
('4. Detectar superficies horizontales del entorno.', 'Probar con luz uniforme; no usar superficies de vidrio, espejo o brillo intenso.'),
('5. Mostrar instrucciones mientras se escanea el entorno.', 'Usar mensajes legibles y no bloquear la vista de camara.'),
('6. Recibir un toque sobre un plano horizontal para fijar la zona.', 'Validar el punto dentro del plano detectado; no aceptar toques sobre botones.'),
('7. Dibujar el marcador circular de la zona seleccionada.', 'Mantener contraste visual y evitar que el marcador oculte la escena.'),
('8. Habilitar Jugar solo cuando exista una zona valida.', 'Prevenir inicio sin superficie; informar al usuario que puede mover la zona.'),
('9. Crear y reutilizar hoyos y topos para la partida.', 'Evitar creacion excesiva de objetos; liberar recursos al terminar.'),
('10. Calcular posiciones dentro del plano y alejadas del borde.', 'Respetar margen de seguridad y separacion entre objetivos.'),
('11. Mostrar tipos de topo con puntos diferenciados.', 'Mantener colores distinguibles y textos comprensibles.'),
('12. Activar la partida de 60 segundos y el temporizador.', 'Comprobar que el tiempo no continúe despues del cierre.'),
('13. Detectar el toque dirigido a un topo visible.', 'Usar capa de colision especifica; descartar toques de interfaz.'),
('14. Registrar el acierto, la puntuacion y el combo.', 'Validar un solo registro por objetivo; evitar duplicacion de eventos.'),
('15. Mostrar la animacion de martillo en el punto de impacto.', 'Mantener la animacion breve para no afectar rendimiento ni visibilidad.'),
('16. Ocultar el topo no golpeado y reiniciar la racha.', 'Aplicar reglas consistentes y comunicar el resultado en la interfaz.'),
('17. Redistribuir el objetivo en una nueva posicion valida.', 'No ubicar objetivos fuera de la superficie ni demasiado juntos.'),
('18. Mostrar pantalla de fin con puntuacion y mejor combo.', 'Desactivar la entrada de golpes para evitar acciones posteriores.'),
('19. Permitir reiniciar o cambiar la zona de juego.', 'Restablecer variables y objetos antes de una nueva partida.'),
('20. Probar deteccion, colocacion, golpes, interfaz y cierre.', 'Registrar incidencias; proteger el equipo movil y cuidar la bateria.'),
('21. Documentar resultados con capturas verificables.', 'Usar solo evidencias reales; respetar propiedad intelectual de recursos.')]
for r, (a,b) in enumerate(steps, 1):
    set_cell(ops.cell(r,0), a, 7.5); set_cell(ops.cell(r,1), b, 7.5)

# Esquema de propuesta.
scheme = doc.tables[8]
set_cell(scheme.cell(1,1), 'ESCANEAR SUPERFICIE → FIJAR ZONA → INICIAR PARTIDA → TOCAR TOPOS → MOSTRAR RESULTADO', 9, True)
set_cell(scheme.cell(2,1), 'Golpea al Topo en realidad aumentada', 8)
set_cell(scheme.cell(2,2), 'Interaccion tactil y control visual', 8)

# Control y valoración. La tabla de control tiene celdas combinadas en el archivo
# oficial; se conserva para que la validación final la complete el estudiante o docente.
score = doc.tables[10]
for r in range(1,6): set_cell(score.cell(r,3), '[POR COMPLETAR]', 8)

# Pie de esquema y campos repetidos.
for row in doc.tables[8].rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            if '[NOMBRE DEL TEMA DEL TRABAJO FINAL]' in p.text:
                p.text = 'Videojuego de realidad aumentada Golpea al Topo'
            elif '[APELLIDOS Y NOMBRES]' in p.text:
                p.text = '[POR COMPLETAR: APELLIDOS Y NOMBRES]'

doc.core_properties.title = 'Trabajo Final Videojuego de realidad aumentada Golpea al Topo'
doc.core_properties.subject = 'Propuesta y desarrollo de videojuego de realidad aumentada'
doc.core_properties.author = ''
doc.save(OUT)
print(OUT)
